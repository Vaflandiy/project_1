from openpyxl import load_workbook, Workbook
from openpyxl.chart import BarChart, Reference
from random import choice, uniform, randint

wb = load_workbook("/users/denis/PycharmProjects/pythonProject/example.xlsx")

sheet = wb.active

string = 'AaBbCcDdEeFfGgHhIiJjKkLlMmNnOoPpQqRrSsTtYuVvWwXxYyZz1234567890'

for i in range(2, 1002):
    code = ''
    target_num_str = 'A' + str(i)
    target_8_sim = 'B' + str(i)
    target_float_sim = 'C' + str(i)
    target_number_from_neg_to_pos = 'D' + str(i)
    target_data_in_random_year_february = 'E' + str(i)
    for k in range(8):
        char = choice(string)
        code = code + char
    random_float_number = uniform(-1, 1)
    number_from_neg_to_pos = randint(-100, 100)
    year = randint(0, 2023)
    if year % 4 == 0:
        day = randint(1, 29)
    else:
        day = randint(1, 28)
    random_date = (str(day) + '.02.' + str(year))
    print(code)
    sheet[target_num_str] = str(i)
    sheet[target_8_sim] = code
    sheet[target_float_sim] = random_float_number
    sheet[target_number_from_neg_to_pos] = number_from_neg_to_pos
    sheet[target_data_in_random_year_february] = random_date

sheet['A1'] = 'Номер'
sheet['B1'] = '8 Символов'
sheet['C1'] = 'Число-дробь'
sheet['D1'] = 'Случайное число'
sheet['E1'] = 'Случайная дата'

chart = BarChart()
chart.title = 'Заголовок'
data = Reference(sheet, min_col=4, min_row=1, max_col=4, max_row=1002)

chart.add_data(data, titles_from_data=True)
sheet.add_chart(chart, 'G2')

wb.save("/users/denis/PycharmProjects/pythonProject/training.xlsx")
